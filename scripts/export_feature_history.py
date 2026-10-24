"""Build the reviewable iteration workbook from docs/iterations/feature-history.json.
Requires openpyxl 3.1+: python3 -m pip install openpyxl
Run: python3 scripts/export_feature_history.py
"""
import json
from collections import Counter
from datetime import date
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'docs/iterations/feature-history.json'
OUTPUT = ROOT / 'deliverables/qingbao-feature-history.xlsx'
BASE = 'https://github.com/Jameyan42/-'
NAVY, TEAL, GRAY = '183153', '087F8C', '607086'


def make_sheet(wb, name, title, note, summary, headers, widths, rows, table_name, height=84):
    ws = wb.create_sheet(name)
    n = len(headers)
    for row, text, color, size in [(1, title, NAVY, 20), (2, note, GRAY, 10), (3, summary, TEAL, 11)]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
        cell = ws.cell(row, 1, text)
        cell.font = Font(name='微软雅黑', size=size, color=color, bold=row != 2)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 38
    ws.row_dimensions[3].height = 25
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.freeze_panes = 'E6' if n > 10 else 'B6'
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(5, col, header)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[5].height = 30
    for idx, values in enumerate(rows, 6):
        ws.row_dimensions[idx].height = height
        for col, value in enumerate(values, 1):
            cell = ws.cell(idx, col, value)
            cell.font = Font(name='微软雅黑', size=11, color=NAVY)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=Side(style='hair', color='DFE7EF'))
            if isinstance(value, date):
                cell.number_format = 'yyyy-mm-dd'
            if isinstance(value, str) and value.startswith('https://'):
                cell.hyperlink = value
                cell.font = Font(name='微软雅黑', size=10, color='1264A3', underline='single')
            if value in ['已完成', '建议下一阶段', '待排期', '待用户参与', '暂缓']:
                color = 'DFF2EC' if value == '已完成' else 'FFF0CE'
                cell.fill = PatternFill('solid', fgColor=color)
                cell.font = Font(name='微软雅黑', size=11, bold=True, color=TEAL if value == '已完成' else '8A5B09')
    table = Table(displayName=table_name, ref=f'A5:{get_column_letter(n)}{5+len(rows)}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = TEAL
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:5'
    ws.print_options.horizontalCentered = True
    ws.print_area = table.ref
    ws.oddFooter.center.text = '轻报 · 功能迭代记录'
    ws.oddFooter.right.text = '第 &P 页'
    return ws


def main():
    data = json.loads(SOURCE.read_text())
    history, stages = data['history'], data['stages']
    assert [r['sequence'] for r in history] == list(range(1, len(history)+1))
    assert all(r['status'] == '已完成' for r in history)
    assert len({r['version'] for r in stages}) == len(stages)
    stage_names = {s['version'] for s in stages}
    assert all(r['stage'] in stage_names for r in history)
    assert all(date.fromisoformat(r['date']) for r in history)
    for row in history:
        assert (ROOT / row['evidence']).exists(), row['evidence']
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = '轻报功能迭代记录'
    wb.properties.subject = '产品经理面试作品：历史交付、演示入口与下一步计划'
    wb.properties.creator = '轻报项目'
    counts = Counter(r['stage'] for r in history)
    make_sheet(wb, '阶段总览', '轻报｜功能迭代记录',
        '依据 Git 提交与阶段验收回溯；日期沿用阶段验收口径，v0.21起按上海日期记录；不是每项功能的独立完成时间。',
        f"更新：{data['updated']}  |  已交付 {len(stages)} 个阶段、{len(history)} 项内容  |  当前功能版本 {stages[-1]["version"]}  |  后续计划不计入已完成",
        ['阶段顺序','版本/阶段','提交日期','本阶段主题','新增项数','状态','验证记录','源码证据'],
        [12,16,15,30,12,14,56,46],
        [[int(s['order']),s['version'],date.fromisoformat(s['date']),s['title'],counts[s['version']],'已完成',s['verification'],f"{BASE}/commit/{s['commit']}"] for s in stages], 'StageOverview', 74)
    make_sheet(wb, '历史新增明细', '轻报｜历史新增内容（按阶段顺序）',
        '同阶段各项按演示逻辑整理，未宣称逐项开发先后。初版与功能增强无可确认版本号，不补造 v0.1～v0.4。可用表头筛选版本和模块。',
        '阅读顺序：序号 → 新增内容 → 解决的问题 → 页面入口 → 演示方法；可横向滚动查看验证与边界。',
        ['序号','版本/阶段','提交日期','模块','新增内容','状态','解决的问题/价值','页面入口','建议演示方法','验证依据','当前边界/说明','阶段/提交证据'],
        [9,15,15,18,32,13,36,33,40,43,43,48],
        [[r['sequence'],r['stage'],date.fromisoformat(r['date']),r['module'],r['feature'],r['status'],r['value'],r['entry'],r['demo'],r['verification'],r['limits'],
          f"{BASE}/" + (f"blob/{r['commit']}/README.md" if r['evidence']=='README.md' else f"blob/main/{r['evidence']}")] for r in history], 'FeatureHistory', 88)
    make_sheet(wb, '下一步计划', '轻报｜下一步计划',
        '以下是建议及待办，尚未交付。已完成项已移入历史新增明细，不继续混在计划中。',
        '建议顺序：今天先亲自演练与准备答辩 → 真实可用性验证；费用分摊、多法人及制度扩展按反馈排期；规则预审已交付，后续按使用反馈扩展；真实AI按用户选择暂缓。',
        ['顺序','优先级','计划内容','状态','要解决的问题','建议范围','完成验收标准','依赖与说明'],
        [10,12,28,19,40,48,54,46],
        [[r[k] for k in ['order','priority','name','status','problem','scope','acceptance','dependency']] for r in data['roadmap']], 'Roadmap', 100)
    notes = [
        ['如何打开','使用 Microsoft Excel、WPS 或 LibreOffice 打开本文件；这是真正的 .xlsx 工作簿。'],
        ['如何查历史','先看阶段总览，再到历史新增明细按版本或模块筛选；表头和前四列已冻结。'],
        ['顺序与日期','按 Git 功能提交顺序分阶段；同阶段按功能/演示逻辑编号。早期日期沿用源码提交记录，v0.21起按上海日期记录；不将阶段日期当作每项独立完成或上线时间。'],
        ['已完成的含义','已存在于代码和交付记录；每条为阶段新增或改进，不等于独立模块数量，也不代表真实客户验证、生产安全审计或业务收益。'],
        ['测试数据口径','测试数量为当时整个阶段的记录，不是每一项功能的独立测试数量；v0.6、v0.7 含整套测试与补充场景，详见对应验收。'],
        ['线上验证边界','最近版本 Pages 已部署并核对在线产物；浏览器交互在本地同构生产产物验证。云端公网证书链限制未绕过。'],
        ['原型能力边界','数据仅在当前浏览器；身份选择是模拟机制，付款不转账；OCR 不等于验真。请使用虚构资料。'],
        ['后续维护规则','每个完整阶段结束，先更新 JSON 源记录，再重新导出本 Excel；新增记录追加序号，不将计划提前标为已完成。'],
        ['可持续维护来源','docs/iterations/feature-history.json；导出命令：python3 scripts/export_feature_history.py（依赖 openpyxl 3.1+）。'],
        ['与需求文档区别','Excel 追踪“何时新增了什么”；需求及设计文档解释业务规则；阶段验收记录测试证据，三者同时维护。'],
        ['本轮更新','R001 v1仅交付独立住宿费审核、15条城市职级制度、整数分/日历、原始结果与独立人工复核及历史标准快照；不扩展其他业务，无真实企业收益或生产鉴权。'],
        ['在线演示入口','https://jameyan42.github.io/-/?v=0.26'],
    ]
    make_sheet(wb, '使用与维护', '轻报｜使用与维护说明',
        '用于作品进度回顾、面试准备和后续连续迭代；不编造真实用户、访谈或业务提升数字。',
        f"记录口径更新于 {data['updated']}", ['项目','说明'], [25,120], notes, 'WorkbookNotes', 50)
    OUTPUT.parent.mkdir(exist_ok=True)
    wb.save(OUTPUT)
    verify = load_workbook(OUTPUT)
    assert verify.sheetnames == ['阶段总览','历史新增明细','下一步计划','使用与维护']
    assert verify['历史新增明细'].max_row == 5 + len(history)
    assert verify['下一步计划'].max_row == 5 + len(data['roadmap'])
    assert sum(counts.values()) == len(history)
    for ws in verify:
        assert ws.freeze_panes and ws.tables
        assert ws.cell(5,1).value
    print(f'Created and reopened {OUTPUT.name}: {len(stages)} stages, {len(history)} completed features, {len(data["roadmap"])} plans, {len(verify.sheetnames)} sheets.')


if __name__ == '__main__':
    main()
